"""Parse AdControl 'Informe avanzado de cierres de caja' Excel exports."""

from __future__ import annotations

import io
import logging
import re
import unicodedata
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .excel_updater import DailySaleRow

# Minimal OOXML stylesheet — used when AdControl/LibreOffice exports ship broken fills
# that make openpyxl raise: TypeError: Fill() takes no arguments / expected Fill.
_MINIMAL_STYLES_XML = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="1"><font><sz val="11"/><color theme="1"/><name val="Calibri"/><family val="2"/></font></fonts>
  <fills count="2">
    <fill><patternFill patternType="none"/></fill>
    <fill><patternFill patternType="gray125"/></fill>
  </fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>
  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>
"""


def _is_stylesheet_error(exc: BaseException) -> bool:
    text = str(exc)
    return (
        "Fill" in text
        or "fill" in text.lower()
        or "Style" in type(exc).__name__
        or "stylesheet" in text.lower()
    )


def _workbook_bytes_without_styles(path: Path) -> io.BytesIO:
    buf = io.BytesIO()
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename.replace("\\", "/")
            if name == "xl/styles.xml" or name.endswith("/styles.xml"):
                data = _MINIMAL_STYLES_XML
            zout.writestr(info, data)
    buf.seek(0)
    return buf


def load_cierres_workbook(path: str | Path, **kwargs: Any):
    """load_workbook that survives broken styles in some AdControl .xlsx exports."""
    path = Path(path)
    kwargs.setdefault("data_only", True)
    kwargs.setdefault("read_only", True)
    try:
        return load_workbook(path, **kwargs)
    except TypeError as exc:
        if not _is_stylesheet_error(exc):
            raise
        logging.warning(
            "Workbook %s has invalid Excel styles (%s); reloading without styles",
            path.name,
            exc,
        )
        return load_workbook(_workbook_bytes_without_styles(path), **kwargs)
    except ValueError as exc:
        # Newer openpyxl wraps stylesheet failures
        if not _is_stylesheet_error(exc):
            raise
        logging.warning(
            "Workbook %s has invalid Excel styles (%s); reloading without styles",
            path.name,
            exc,
        )
        return load_workbook(_workbook_bytes_without_styles(path), **kwargs)

# Flexible header aliases (export typos / renames).
_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "ubicacion": ("ubicacion", "ubicacion comercial"),
    "estado": ("estado",),
    "cantidad_cierre": ("cantidad de cierre",),
    "fondo_caja": ("fondo de caja",),
    # Sales cash for Ventas Diarias EFECTIVO (excludes fondo de caja float)
    "pago_efectivo": (
        "pago en efectivo",
        "pago en efectivos",
        "pagos en efectivo",
        "total pago en efectivo",
        "total en pago en efectivo",
    ),
    "hora_apertura": ("hora de apertura",),
    "tarjeta": ("total en pago con tarjeta", "pago con tarjeta"),
    "transferencia": (
        "transeferencia bancaria",  # AdControl typo
        "transferencia bancaria",
    ),
    "pedidos_ya": ("total en pedidosya", "pedidosya", "pedidos ya"),
    "otros_pagos": ("total en otros pagos", "otros pagos"),
}

# Expected float per store when export omits Fondo de Caja (fallback only).
_DEFAULT_FONDO_BY_UBICACION: dict[str, float] = {
    "La Cata SJM": 5000.0,
    "La Cata DF": 2500.0,
}

_CLOSED_ESTADOS = {"", "close", "cerrado", "closed", "cierre"}


def _strip_accents(text: str) -> str:
    return "".join(
        ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn"
    )


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    text = _strip_accents(str(value)).lower().strip()
    return " ".join(text.split())


def _parse_money(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    # Prefer last money-looking token (handles "RD$ 33,597.60")
    matches = re.findall(r"[-0-9,]+\.\d{2}", text)
    if matches:
        return float(matches[-1].replace(",", ""))
    cleaned = (
        text.replace("RD$", "")
        .replace("rd$", "")
        .replace(",", "")
        .replace(" ", "")
        .strip()
    )
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError as exc:
        raise ValueError(f"Could not parse money value: {value!r}") from exc


def _parse_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)):
        from openpyxl.utils.datetime import from_excel

        return from_excel(value)
    text = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"Could not parse datetime: {value!r}")


def _map_ubicacion(raw: Any) -> str:
    compact = _strip_accents(str(raw or "")).upper()
    compact_nospace = re.sub(r"\s+", "", compact)

    has_sjm = bool(
        re.search(r"\(\s*SAN\s*JUAN\s*\)", compact)
        or re.search(r"\bSAN\s*JUAN\b", compact)
        or re.search(r"\bSJM\b", compact)
        or "SANJUAN" in compact_nospace
    )
    has_df = bool(
        re.search(r"\(\s*DEFILLO\s*\)", compact)
        or re.search(r"\bDEFILLO\b", compact)
        or re.search(r"\bDF\b", compact)
        or "DEFILLO" in compact_nospace
    )
    # Avoid matching "DF" inside unrelated words when both markers absent —
    # require DEFILLO or explicit (DEFILLO) / standalone DF token (already above).

    if has_sjm and has_df:
        raise ValueError(f"Ambiguous ubicación in cierres export: {raw!r}")
    if has_sjm:
        return "La Cata SJM"
    if has_df:
        return "La Cata DF"
    raise ValueError(f"Unknown ubicación in cierres export: {raw!r}")


def _resolve_headers(headers: list[Any]) -> dict[str, int]:
    normed = [_norm_header(h) for h in headers]
    resolved: dict[str, int] = {}
    for key, aliases in _HEADER_ALIASES.items():
        for idx, name in enumerate(normed):
            if name in aliases:
                resolved[key] = idx
                break
    required = (
        "ubicacion",
        "hora_apertura",
        "tarjeta",
        "transferencia",
        "otros_pagos",
    )
    missing = [k for k in required if k not in resolved]
    # Need either Pago en efectivo, or Cantidad de cierre (+ optional Fondo) to derive it.
    if "pago_efectivo" not in resolved and "cantidad_cierre" not in resolved:
        missing.append("pago_efectivo|cantidad_cierre")
    if missing:
        raise ValueError(
            "Cierres Excel missing required columns: "
            + ", ".join(missing)
            + f". Found headers: {headers!r}"
        )
    return resolved


@dataclass
class CierreShift:
    fecha: date
    ubicacion: str
    efectivo: float
    tarjeta: float
    transferencias: float
    pedidos_ya: float
    fondo_caja: float = 0.0
    cantidad_cierre: float | None = None
    usuario: str | None = None
    estado: str | None = None
    sheet: str | None = None


@dataclass
class CierresParseStats:
    sheets_read: list[str] = field(default_factory=list)
    sheets_skipped: list[str] = field(default_factory=list)
    raw_ubicaciones: dict[str, int] = field(default_factory=dict)
    skipped_open: int = 0
    skipped_unknown_ubicacion: dict[str, int] = field(default_factory=dict)
    skipped_errors: int = 0
    rows_seen: int = 0
    headers_resolved: list[str] = field(default_factory=list)
    efectivo_source: str | None = None
    sample_breakdown: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "sheets_read": list(self.sheets_read),
            "sheets_skipped": list(self.sheets_skipped),
            "raw_ubicaciones": dict(self.raw_ubicaciones),
            "skipped_open": self.skipped_open,
            "skipped_unknown_ubicacion": dict(self.skipped_unknown_ubicacion),
            "skipped_errors": self.skipped_errors,
            "rows_seen": self.rows_seen,
            "headers_resolved": list(self.headers_resolved),
            "efectivo_source": self.efectivo_source,
            "sample_breakdown": list(self.sample_breakdown),
        }


def _efectivo_from_row(
    row: tuple[Any, ...],
    cols: dict[str, int],
    ubicacion: str,
) -> tuple[float, float, float | None, str]:
    """Return (efectivo_ventas, fondo_caja, cantidad_cierre, source_label).

    Ventas Diarias EFECTIVO must be sales cash only:
    prefer AdControl **Pago en efectivo**, else Cantidad de cierre − Fondo de Caja.
    """
    cantidad: float | None = None
    if "cantidad_cierre" in cols:
        cantidad = _parse_money(row[cols["cantidad_cierre"]])

    default_fondo = _DEFAULT_FONDO_BY_UBICACION.get(ubicacion, 0.0)
    if "fondo_caja" in cols:
        fondo = _parse_money(row[cols["fondo_caja"]])
        # Empty/zero fondo in export is not trustworthy for these stores
        if fondo <= 0 and default_fondo > 0:
            logging.warning(
                "%s: Fondo de Caja missing/zero in row; using default %.2f",
                ubicacion,
                default_fondo,
            )
            fondo = default_fondo
    else:
        fondo = default_fondo

    source = "cantidad_cierre_minus_fondo"
    if "pago_efectivo" in cols:
        efectivo = _parse_money(row[cols["pago_efectivo"]])
        source = "pago_en_efectivo"
        # Some exports duplicate drawer total into Pago en efectivo
        if cantidad is not None and fondo > 0 and abs(efectivo - cantidad) < 0.01:
            logging.warning(
                "%s: Pago en efectivo (%.2f) equals Cantidad de cierre; "
                "subtracting fondo %.2f",
                ubicacion,
                efectivo,
                fondo,
            )
            efectivo = max(0.0, cantidad - fondo)
            source = "cantidad_cierre_minus_fondo_guard"
    elif cantidad is not None:
        efectivo = max(0.0, cantidad - fondo)
    else:
        raise ValueError("Cierres row missing Pago en efectivo and Cantidad de cierre")

    if cantidad is not None and abs(efectivo - cantidad) < 0.01 and fondo > 0:
        # Final safety: never write drawer total when float is known
        logging.warning(
            "%s: EFECTIVO still equals Cantidad de cierre (%.2f); forcing minus fondo %.2f",
            ubicacion,
            cantidad,
            fondo,
        )
        efectivo = max(0.0, cantidad - fondo)
        source = "cantidad_cierre_minus_fondo_forced"

    if cantidad is not None and source.startswith("pago") and fondo > 0:
        expected = max(0.0, round(cantidad - fondo, 2))
        if abs(expected - round(efectivo, 2)) > 0.01:
            logging.warning(
                "Cierres mismatch %s: cantidad_cierre=%s fondo=%s "
                "pago_efectivo=%s (expected cantidad-fondo=%s)",
                ubicacion,
                cantidad,
                fondo,
                efectivo,
                expected,
            )
    return efectivo, fondo, cantidad, source


def _parse_sheet(
    ws: Any,
    *,
    sheet_name: str,
    stats: CierresParseStats,
) -> list[CierreShift]:
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        stats.sheets_skipped.append(f"{sheet_name}:empty")
        return []

    try:
        cols = _resolve_headers(list(header_row))
    except ValueError as exc:
        stats.sheets_skipped.append(f"{sheet_name}:{exc}")
        logging.info("Skipping sheet %r (not cierres layout): %s", sheet_name, exc)
        return []

    if not stats.headers_resolved:
        stats.headers_resolved = sorted(cols.keys())
        logging.info(
            "Cierres headers on %r: %s | pago_efectivo_col=%s fondo_col=%s",
            sheet_name,
            stats.headers_resolved,
            "pago_efectivo" in cols,
            "fondo_caja" in cols,
        )

    shifts: list[CierreShift] = []
    for row in rows:
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        stats.rows_seen += 1
        ubic_raw = row[cols["ubicacion"]]
        if ubic_raw is None or str(ubic_raw).strip() == "":
            continue
        raw_label = str(ubic_raw).strip()
        stats.raw_ubicaciones[raw_label] = stats.raw_ubicaciones.get(raw_label, 0) + 1

        estado = row[cols["estado"]] if "estado" in cols else None
        estado_norm = str(estado).strip().lower() if estado is not None else ""
        if estado is not None and estado_norm not in _CLOSED_ESTADOS:
            stats.skipped_open += 1
            continue

        try:
            ubicacion = _map_ubicacion(ubic_raw)
        except ValueError:
            stats.skipped_unknown_ubicacion[raw_label] = (
                stats.skipped_unknown_ubicacion.get(raw_label, 0) + 1
            )
            logging.warning("Skipping unknown ubicación in %s: %r", sheet_name, raw_label)
            continue

        try:
            fecha = _parse_datetime(row[cols["hora_apertura"]]).date()
            transferencia = _parse_money(row[cols["transferencia"]])
            otros = _parse_money(row[cols["otros_pagos"]])
            pedidos = (
                _parse_money(row[cols["pedidos_ya"]]) if "pedidos_ya" in cols else 0.0
            )
            usuario = None
            for idx, h in enumerate(header_row):
                if _norm_header(h) == "usuario":
                    raw_user = row[idx]
                    if raw_user is not None:
                        usuario = str(raw_user).strip() or None
                    break
            efectivo, fondo, cantidad, ef_source = _efectivo_from_row(row, cols, ubicacion)
            if stats.efectivo_source is None:
                stats.efectivo_source = ef_source
            if len(stats.sample_breakdown) < 6:
                stats.sample_breakdown.append(
                    {
                        "fecha": fecha.isoformat(),
                        "ubicacion": ubicacion,
                        "cantidad_cierre": cantidad,
                        "fondo_caja": fondo,
                        "efectivo_escrito": efectivo,
                        "source": ef_source,
                    }
                )
            shifts.append(
                CierreShift(
                    fecha=fecha,
                    ubicacion=ubicacion,
                    efectivo=efectivo,
                    tarjeta=_parse_money(row[cols["tarjeta"]]),
                    transferencias=transferencia + otros,
                    pedidos_ya=pedidos,
                    fondo_caja=fondo,
                    cantidad_cierre=cantidad,
                    usuario=usuario,
                    estado=str(estado).strip() if estado is not None else None,
                    sheet=sheet_name,
                )
            )
        except Exception as exc:  # noqa: BLE001
            stats.skipped_errors += 1
            logging.warning("Skipping bad cierres row in %s: %s", sheet_name, exc)

    stats.sheets_read.append(sheet_name)
    logging.info(
        "Sheet %r: %s closed shifts parsed",
        sheet_name,
        len(shifts),
    )
    return shifts


def parse_cierres_workbook(
    path: str | Path,
    *,
    return_stats: bool = False,
) -> list[CierreShift] | tuple[list[CierreShift], CierresParseStats]:
    """Parse closed cash-register shifts from all sheets in the cierres Excel."""
    path = Path(path)
    name_l = path.name.lower()
    if "registrar" in name_l and "cierres" not in name_l:
        logging.warning(
            "File %r looks like a 'Registrar informe' export, not "
            "'Informe avanzado de cierres de caja'. Wrong report type usually fails "
            "column detection — re-export the advanced cash-closures report.",
            path.name,
        )
    wb = load_cierres_workbook(path, data_only=True, read_only=True)
    stats = CierresParseStats()
    try:
        shifts: list[CierreShift] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            shifts.extend(_parse_sheet(ws, sheet_name=sheet_name, stats=stats))
        if not shifts:
            raise ValueError(
                "No closed cierres rows found in workbook. "
                f"sheets_read={stats.sheets_read} sheets_skipped={stats.sheets_skipped} "
                f"raw_ubicaciones={stats.raw_ubicaciones} "
                f"skipped_open={stats.skipped_open} "
                f"skipped_unknown={stats.skipped_unknown_ubicacion}"
            )
        logging.info(
            "Cierres parse %s: %s shifts from sheets %s | raw ubicaciones=%s | "
            "skipped_open=%s unknown=%s errors=%s",
            path.name,
            len(shifts),
            stats.sheets_read,
            stats.raw_ubicaciones,
            stats.skipped_open,
            stats.skipped_unknown_ubicacion,
            stats.skipped_errors,
        )
        if return_stats:
            return shifts, stats
        return shifts
    finally:
        wb.close()


def aggregate_cierres(shifts: list[CierreShift]) -> list[DailySaleRow]:
    """Sum shifts by (fecha, ubicacion) into Ventas Diarias rows."""
    buckets: dict[tuple[date, str], DailySaleRow] = {}
    for shift in shifts:
        key = (shift.fecha, shift.ubicacion)
        if key not in buckets:
            buckets[key] = DailySaleRow(
                fecha=shift.fecha,
                ubicacion=shift.ubicacion,
                efectivo=shift.efectivo,
                tarjeta_bruta=shift.tarjeta,
                transferencias=shift.transferencias,
                notas_credito=0.0,
                itbs_cobrado=None,  # not present in cierres export
                pedidos_ya=shift.pedidos_ya,
            )
            continue
        existing = buckets[key]
        buckets[key] = DailySaleRow(
            fecha=existing.fecha,
            ubicacion=existing.ubicacion,
            efectivo=existing.efectivo + shift.efectivo,
            tarjeta_bruta=existing.tarjeta_bruta + shift.tarjeta,
            transferencias=existing.transferencias + shift.transferencias,
            notas_credito=existing.notas_credito,
            itbs_cobrado=existing.itbs_cobrado,
            pedidos_ya=(existing.pedidos_ya or 0.0) + shift.pedidos_ya,
        )
    return sorted(buckets.values(), key=lambda r: (r.fecha, r.ubicacion))


def summarize_cierres(rows: list[DailySaleRow]) -> dict[str, Any]:
    by_ubic: dict[str, int] = defaultdict(int)
    for row in rows:
        by_ubic[row.ubicacion] += 1
    return {
        "days_locations": len(rows),
        "by_ubicacion": dict(by_ubic),
        "date_from": rows[0].fecha.isoformat() if rows else None,
        "date_to": rows[-1].fecha.isoformat() if rows else None,
    }


def is_cierres_workbook(path: str | Path) -> bool:
    """True if workbook looks like AdControl 'Informe avanzado de cierres de caja'."""
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    # Ignore Excel lock / temp files
    if path.name.startswith("~$") or path.name.startswith("."):
        return False
    # Fast path: common export filename (spaces, -2 copies, etc. are fine)
    name = path.name.lower()
    if ("cierres" in name and "caja" in name) or (
        "informe" in name and "cierres" in name
    ):
        return True
    try:
        wb = load_cierres_workbook(path, data_only=True, read_only=True)
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                if not header_row:
                    continue
                try:
                    _resolve_headers(list(header_row))
                    return True
                except ValueError:
                    continue
            return False
        finally:
            wb.close()
    except Exception as exc:  # noqa: BLE001
        logging.debug("Not a cierres workbook %s: %s", path.name, exc)
        return False


def resolve_cierres_path(path: str | Path, base_dir: str | Path | None = None) -> Path:
    """Resolve a cierres Excel path; try inbox/ when given a bare filename."""
    raw = Path(path).expanduser()
    candidates: list[Path] = []
    if raw.is_absolute():
        candidates.append(raw)
    else:
        candidates.append(Path.cwd() / raw)
        if base_dir is not None:
            base = Path(base_dir)
            candidates.append(base / raw)
            candidates.append(base / "inbox" / raw.name)
            # Also allow just the filename while cwd is elsewhere
            candidates.append(base / "inbox" / Path(path).name)

    seen: set[Path] = set()
    for cand in candidates:
        try:
            resolved = cand.resolve()
        except OSError:
            continue
        if resolved in seen:
            continue
        seen.add(resolved)
        if resolved.is_file():
            return resolved

    searched = ", ".join(str(c) for c in candidates)
    raise FileNotFoundError(
        f"Cierres Excel not found: {path!r}. Looked in: {searched}. "
        "Put the file in inbox/ or pass a full path in quotes "
        "(spaces in the name require quotes)."
    )

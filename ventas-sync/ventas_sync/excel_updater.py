"""Upsert Ventas Diarias rows in the La Cata workbook."""

from __future__ import annotations

import json
import time
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .pdf_parser import TotalesGenerales


@dataclass
class DailySaleRow:
    fecha: date
    ubicacion: str
    efectivo: float
    tarjeta_bruta: float
    transferencias: float
    notas_credito: float
    itbs_cobrado: float | None = None


def totales_to_row(totales: TotalesGenerales, ubicacion: str) -> DailySaleRow:
    return DailySaleRow(
        fecha=totales.fecha.date(),
        ubicacion=ubicacion,
        efectivo=totales.efectivo,
        tarjeta_bruta=totales.tarjeta,
        transferencias=totales.otros_pagos,
        notas_credito=totales.nota_credito,
        itbs_cobrado=totales.itbis,
    )


def _as_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial fallback via openpyxl already gives datetime usually
        from openpyxl.utils.datetime import from_excel

        return from_excel(value).date()
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _find_row(ws: Worksheet, fecha: date, ubicacion: str) -> int | None:
    for row in range(2, ws.max_row + 1):
        row_date = _as_date(ws.cell(row, 1).value)
        row_ubic = ws.cell(row, 2).value
        if row_date == fecha and row_ubic == ubicacion:
            return row
    return None


def _first_append_row(ws: Worksheet) -> int:
    """First row where FECHA is empty (pre-seeded formula rows are OK to fill)."""
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value is None and ws.cell(row, 2).value is None:
            return row
    return ws.max_row + 1


def _write_formulas(ws: Worksheet, row: int, fee_rate_cell: str) -> None:
    ws.cell(row, 5).value = f"=D{row}*{fee_rate_cell}"  # FEE
    ws.cell(row, 6).value = f"=D{row}-E{row}"  # TARJETA NETA
    ws.cell(row, 10).value = f"=C{row}+F{row}+G{row}-I{row}"  # TOTAL
    ws.cell(row, 13).value = f"=J{row}-L{row}"  # VENTAS SIN ITBIS


def apply_row(ws: Worksheet, row: int, data: DailySaleRow, fee_rate_cell: str) -> None:
    ws.cell(row, 1).value = datetime.combine(data.fecha, datetime.min.time())
    ws.cell(row, 2).value = data.ubicacion
    ws.cell(row, 3).value = data.efectivo
    ws.cell(row, 4).value = data.tarjeta_bruta
    # H PEDIDOS YA left alone if somehow set; default blank on new rows
    if ws.cell(row, 8).value is None:
        ws.cell(row, 8).value = None
    ws.cell(row, 7).value = data.transferencias
    ws.cell(row, 9).value = data.notas_credito
    if data.itbs_cobrado is not None:
        ws.cell(row, 12).value = data.itbs_cobrado
    elif ws.cell(row, 12).value is None:
        ws.cell(row, 12).value = 0
    # Never overwrite deposito (N / col 14) if already filled
    _write_formulas(ws, row, fee_rate_cell)


def upsert_daily_sale(
    excel_path: str | Path,
    sheet_name: str,
    data: DailySaleRow,
    fee_rate_cell: str = "Setup!$B$6",
) -> dict[str, Any]:
    path = Path(excel_path)
    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")
    ws = wb[sheet_name]

    existing = _find_row(ws, data.fecha, data.ubicacion)
    if existing is None:
        row = _first_append_row(ws)
        action = "inserted"
    else:
        row = existing
        action = "updated"

    apply_row(ws, row, data, fee_rate_cell)
    wb.save(path)
    wb.close()
    return {"action": action, "row": row, "fecha": data.fecha.isoformat(), "ubicacion": data.ubicacion}


def upsert_with_retry(
    excel_path: str | Path,
    sheet_name: str,
    data: DailySaleRow,
    fee_rate_cell: str = "Setup!$B$6",
    retries: int = 5,
    retry_seconds: float = 2.0,
    pending_path: str | Path | None = None,
) -> dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            result = upsert_daily_sale(excel_path, sheet_name, data, fee_rate_cell)
            result["attempts"] = attempt
            return result
        except PermissionError as exc:
            last_error = exc
            time.sleep(retry_seconds)
        except OSError as exc:
            # Windows/Mac file lock often surfaces as OSError
            last_error = exc
            time.sleep(retry_seconds)

    if pending_path is not None:
        queue_pending(pending_path, data)
        return {
            "action": "queued",
            "error": str(last_error),
            "fecha": data.fecha.isoformat(),
            "ubicacion": data.ubicacion,
        }
    raise RuntimeError(f"Failed to update Excel after {retries} attempts: {last_error}")


def queue_pending(pending_path: str | Path, data: DailySaleRow) -> None:
    path = Path(pending_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    items: list[dict[str, Any]] = []
    if path.exists():
        items = json.loads(path.read_text(encoding="utf-8"))
    payload = asdict(data)
    payload["fecha"] = data.fecha.isoformat()
    items.append(payload)
    path.write_text(json.dumps(items, indent=2), encoding="utf-8")


def flush_pending(
    pending_path: str | Path,
    excel_path: str | Path,
    sheet_name: str,
    fee_rate_cell: str = "Setup!$B$6",
    retries: int = 5,
    retry_seconds: float = 2.0,
) -> list[dict[str, Any]]:
    path = Path(pending_path)
    if not path.exists():
        return []
    items = json.loads(path.read_text(encoding="utf-8"))
    remaining: list[dict[str, Any]] = []
    results: list[dict[str, Any]] = []
    for item in items:
        data = DailySaleRow(
            fecha=date.fromisoformat(item["fecha"]),
            ubicacion=item["ubicacion"],
            efectivo=item["efectivo"],
            tarjeta_bruta=item["tarjeta_bruta"],
            transferencias=item["transferencias"],
            notas_credito=item["notas_credito"],
            itbs_cobrado=item.get("itbs_cobrado"),
        )
        result = upsert_with_retry(
            excel_path,
            sheet_name,
            data,
            fee_rate_cell=fee_rate_cell,
            retries=retries,
            retry_seconds=retry_seconds,
            pending_path=None,
        )
        if result.get("action") == "queued":
            remaining.append(item)
        else:
            results.append(result)
    if remaining:
        path.write_text(json.dumps(remaining, indent=2), encoding="utf-8")
    else:
        path.unlink(missing_ok=True)
    return results
